"""Atualiza a cópia local da aba BASE_2026, sem incluir as abas de contatos."""
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
workbook = load_workbook(ROOT / 'Leilão - 2026.xlsx', data_only=True, read_only=True)
sheet = workbook['BASE_2026']
values = sheet.iter_rows(values_only=True)
headers = [str(value or '').strip() for value in next(values)]
records = []
for values_row in values:
    record = {}
    for key, value in zip(headers, values_row):
        if isinstance(value, (datetime, date)):
            value = value.strftime('%d/%m/%Y')
        record[key] = value.strip() if isinstance(value, str) else value if value is not None else ''
    if record.get('PLACA'):
        records.append(record)
output = ROOT / 'Dashboard_Inteligencia_Logistica/assets/js/workbook-data.js'
output.write_text('// Cópia local da aba BASE_2026. Atualize com scripts/extract_base.py.\nconst workbookBase = ' + json.dumps(records, ensure_ascii=False, indent=2) + ';\n', encoding='utf-8')
workbook.close()
print(f'{len(records)} registros exportados de BASE_2026.')
